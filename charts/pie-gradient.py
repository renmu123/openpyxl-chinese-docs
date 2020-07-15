from openpyxl import Workbook

from openpyxl.chart import (
    PieChart,
    Reference
)
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.fill import (
    GradientFillProperties,
    GradientStop,
    LinearShadeProperties
)
from openpyxl.drawing.colors import SchemeColor

data = [
    ['Pie', 'Sold'],
    ['Apple', 50],
    ['Cherry', 30],
    ['Pumpkin', 10],
    ['Chocolate', 40],
]

wb = Workbook()
ws = wb.active

for row in data:
    ws.append(row)

pie = PieChart()
labels = Reference(ws, min_col=1, min_row=2, max_row=5)
data = Reference(ws, min_col=2, min_row=1, max_row=5)
pie.add_data(data, titles_from_data=True)
pie.set_categories(labels)
pie.title = "Pies sold by category"

# Cut the first slice out of the pie and apply a gradient to it
slice = DataPoint(
    idx=0,
    explosion=20,
    spPr=GraphicalProperties(
        gradFill=GradientFillProperties(
            gsLst=(
                GradientStop(
                    pos=0,
                    prstClr='blue'
                ),
                GradientStop(
                    pos=100000,
                    schemeClr=SchemeColor(
                        val='accent1',
                        lumMod=30000,
                        lumOff=70000
                    )
                )
            )
        )
    )
)
pie.series[0].data_points = [slice]

ws.add_chart(pie, "D1")

wb.save("pie.xlsx")
